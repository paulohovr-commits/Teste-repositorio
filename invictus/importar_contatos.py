#!/usr/bin/env python3
"""
Importa uma lista própria de contatos (WhatsApp) para o CRM, sem custo.

Uso:
    python importar_contatos.py contatos.csv            # só organiza (gera leads_merged.csv)
    python importar_contatos.py contatos.xlsx --montar  # organiza e já gera o index.html

Aceita .csv, .txt (um número por linha) e .xlsx. A coluna do telefone é
detectada sozinha; a do nome é a primeira coluna de texto que não for
telefone. Números repetidos ou inválidos são descartados.

Cada número vira um lead na coluna "Novo" do CRM, com o botão de WhatsApp,
a região pelo DDD e a marcação celular/fixo. Nada é enviado: o CRM só abre
a conversa para você escrever e mandar.
"""
from __future__ import annotations

import csv
import io
import os
import re
import subprocess
import sys
from pathlib import Path

import recorte

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

OUT_CSV = recorte.BASE / "leads_merged.csv"

REGIAO_DDD_SP = {
    "11": "São Paulo e Grande SP", "12": "Vale do Paraíba e Litoral Norte", "13": "Baixada Santista",
    "14": "Bauru e Marília", "15": "Sorocaba e região", "16": "Ribeirão Preto e Franca",
    "17": "São José do Rio Preto", "18": "Presidente Prudente e Araçatuba", "19": "Campinas e Piracicaba",
}
UF_DDD = {
    **{d: "SP" for d in REGIAO_DDD_SP},
    "21": "RJ", "22": "RJ", "24": "RJ", "27": "ES", "28": "ES",
    "31": "MG", "32": "MG", "33": "MG", "34": "MG", "35": "MG", "37": "MG", "38": "MG",
    "41": "PR", "42": "PR", "43": "PR", "44": "PR", "45": "PR", "46": "PR",
    "47": "SC", "48": "SC", "49": "SC", "51": "RS", "53": "RS", "54": "RS", "55": "RS",
    "61": "DF", "62": "GO", "64": "GO", "63": "TO", "65": "MT", "66": "MT", "67": "MS",
    "68": "AC", "69": "RO", "71": "BA", "73": "BA", "74": "BA", "75": "BA", "77": "BA",
    "79": "SE", "81": "PE", "87": "PE", "82": "AL", "83": "PB", "84": "RN", "85": "CE", "88": "CE",
    "86": "PI", "89": "PI", "91": "PA", "93": "PA", "94": "PA", "92": "AM", "97": "AM",
    "95": "RR", "96": "AP", "98": "MA", "99": "MA",
}


def normalizar(bruto: str) -> str:
    """Devolve DDD + número (10 ou 11 dígitos) ou '' se não for telefone brasileiro."""
    dig = re.sub(r"\D", "", bruto or "")
    if dig.startswith("55") and len(dig) in (12, 13):
        dig = dig[2:]
    dig = dig.lstrip("0")
    if len(dig) not in (10, 11) or dig[:2] not in UF_DDD:
        return ""
    if len(dig) == 11 and dig[2] != "9":
        return ""
    return dig


def formatar(dig: str) -> str:
    return f"({dig[:2]}) {dig[2:-4]}-{dig[-4:]}"


def ler_linhas(path: Path) -> list[list[str]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        return [["" if c is None else str(c) for c in row] for row in wb.active.iter_rows(values_only=True)]
    texto = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        return list(csv.reader(io.StringIO(texto), dialeto))
    except csv.Error:
        return [[linha] for linha in texto.splitlines()]


def detectar_colunas(linhas: list[list[str]]) -> tuple[int, int | None]:
    amostra = linhas[:200]
    largura = max((len(r) for r in amostra), default=0)
    tel_hits = [sum(1 for r in amostra if i < len(r) and normalizar(r[i])) for i in range(largura)]
    col_tel = max(range(largura), key=lambda i: tel_hits[i])
    col_nome = None
    for i in range(largura):
        if i == col_tel:
            continue
        textos = sum(1 for r in amostra if i < len(r) and re.search(r"[A-Za-zÀ-ú]{3}", r[i] or ""))
        if textos >= len(amostra) * 0.5:
            col_nome = i
            break
    return col_tel, col_nome


def main(args: list[str]) -> int:
    arquivos = [a for a in args if not a.startswith("--")]
    if len(arquivos) != 1:
        print(__doc__)
        return 1
    path = Path(arquivos[0])
    linhas = ler_linhas(path)
    if not linhas:
        print("[erro] arquivo vazio")
        return 1
    col_tel, col_nome = detectar_colunas(linhas)
    print(f"Coluna do telefone: {col_tel + 1} | coluna do nome: {'nenhuma' if col_nome is None else col_nome + 1}")

    vistos, leads, invalidos = set(), [], 0
    for r in linhas:
        dig = normalizar(r[col_tel]) if col_tel < len(r) else ""
        if not dig:
            invalidos += 1
            continue
        if dig in vistos:
            continue
        vistos.add(dig)
        ddd = dig[:2]
        celular = len(dig) == 11
        uf = UF_DDD[ddd]
        regiao = REGIAO_DDD_SP.get(ddd, f"DDD {ddd}")
        nome = (r[col_nome].strip() if col_nome is not None and col_nome < len(r) else "") or f"Contato {formatar(dig)}"
        leads.append({
            "nome": nome,
            "categoria": "Lista própria",
            "endereco": f"DDD {ddd}, {regiao} - {uf}",
            "telefone": formatar(dig),
            "whatsapp": "55" + dig,
            "celular": "Sim" if celular else "Não",
            "whatsapp_status": "provavel" if celular else "fixo",
            "site": "", "nota": "", "avaliacoes": "", "fotos": "", "recencia_dias": "",
            "score": 50, "nrl": "", "maps": "",
            "query_origem": f"lista_propria:{path.name}",
        })

    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(leads[0].keys()) if leads else ["nome"])
        w.writeheader()
        w.writerows(leads)

    por_uf: dict[str, int] = {}
    for lead in leads:
        uf = lead["endereco"].rsplit("- ", 1)[-1]
        por_uf[uf] = por_uf.get(uf, 0) + 1
    print(f"Contatos válidos e únicos: {len(leads)} | descartados (inválidos/cabeçalho): {invalidos}")
    print(f"Celulares: {sum(1 for x in leads if x['celular'] == 'Sim')} | fixos: {sum(1 for x in leads if x['celular'] == 'Não')}")
    print("Por estado: " + ", ".join(f"{k} {v}" for k, v in sorted(por_uf.items(), key=lambda kv: -kv[1])))
    print(f"Salvo: {OUT_CSV}")

    if "--montar" in args and leads:
        env = {**os.environ, "LEADS_ALVO": str(len(leads))}
        for script in ("consolidate_v2.py", "build_html_v2.py"):
            subprocess.run([sys.executable, str(recorte.BASE / script)], check=True, env=env, cwd=recorte.BASE)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
